# -*- coding: utf-8 -*-
"""
Dosya dönüştürme yardımcıları.
Web katmanı (Flask) burayı çağırır; iş kuralları bu modülde kalır.
"""

from __future__ import annotations

import csv
import io
import json
from pathlib import PurePath

from PIL import Image


def csv_bytes_to_json_bytes(ham_csv: bytes) -> bytes:
    """
    CSV içeriğini UTF-8 (BOM destekli) okuyup JSON'a çevirir.

    İlk satırda anlamlı başlıklar varsa her satır bir nesne (DictReader).
    Aksi halde satırlar dizi olarak kalır (iç içe listeler).
    """
    metin = ham_csv.decode("utf-8-sig")
    okuyucu = csv.reader(io.StringIO(metin))
    satirlar = list(okuyucu)

    if not satirlar:
        veri: list | list[dict] = []
    else:
        baslik = satirlar[0]
        baslik_gecerli = len(baslik) > 0 and all(h.strip() != "" for h in baslik)
        if baslik_gecerli and len(satirlar) > 1:
            veri = [dict(zip(baslik, satir)) for satir in satirlar[1:]]
        else:
            veri = satirlar

    return json.dumps(veri, ensure_ascii=False, indent=2).encode("utf-8")


def csv_bytes_to_tsv_bytes(ham_csv: bytes) -> bytes:
    """CSV baytını UTF-8 okuyup sekme ile ayrılmış metne (TSV) çevirir."""
    metin = ham_csv.decode("utf-8-sig")
    okuyucu = csv.reader(io.StringIO(metin))
    cikis = io.StringIO()
    yazici = csv.writer(cikis, delimiter="\t", lineterminator="\n")
    for satir in okuyucu:
        yazici.writerow(satir)
    return cikis.getvalue().encode("utf-8")


def json_bytes_prettify(ham_json: bytes) -> bytes:
    """JSON metnini doğrular ve okunaklı (girintili) UTF-8 JSON olarak döndürür."""
    metin = ham_json.decode("utf-8-sig")
    veri = json.loads(metin)
    return json.dumps(veri, ensure_ascii=False, indent=2).encode("utf-8")


def json_bytes_to_csv_bytes(ham_json: bytes) -> bytes:
    """
    JSON dizisini (yalnızca nesne listesi, tümünde ortak anahtarlar) CSV'ye çevirir.
    """
    metin = ham_json.decode("utf-8-sig")
    veri = json.loads(metin)
    if not isinstance(veri, list) or len(veri) == 0:
        raise ValueError("json_csv_sadece_dizi")
    if not all(isinstance(o, dict) for o in veri):
        raise ValueError("json_csv_sadece_nesneler")
    anahtarlar = list(veri[0].keys())
    if not anahtarlar:
        raise ValueError("json_csv_bos_nesne")
    for o in veri[1:]:
        if list(o.keys()) != anahtarlar:
            raise ValueError("json_csv_uyumsuz_anahtarlar")

    cikis = io.StringIO()
    yazici = csv.DictWriter(cikis, fieldnames=anahtarlar, extrasaction="ignore", lineterminator="\n")
    yazici.writeheader()
    yazici.writerows(veri)
    return cikis.getvalue().encode("utf-8")


def csv_cikis_json_to_csv_ad(orijinal_ad: str) -> str:
    """JSON kaynağından üretilen CSV adı."""
    kok = PurePath(orijinal_ad or "cikti").stem
    return f"{kok}.csv"


def pdf_bytes_to_txt(ham_pdf: bytes) -> bytes:
    """PDF içinden düz metin çıkarır."""
    from pypdf import PdfReader

    okuyucu = PdfReader(io.BytesIO(ham_pdf))
    parcalar: list[str] = []
    for sayfa in okuyucu.pages:
        parcalar.append(sayfa.extract_text() or "")
    return "\n".join(parcalar).encode("utf-8")


def pdf_bytes_to_ocr_txt(ham_pdf: bytes) -> bytes:
    """PDF sayfalarını resme dönüştürüp OCR (pytesseract) ile metin çıkarır."""
    try:
        from pdf2image import convert_from_bytes
        import pytesseract
    except ImportError:
        raise RuntimeError("PDF OCR için 'pdf2image' ve 'pytesseract' paketleri gereklidir.") from None

    try:
        sayfalar = convert_from_bytes(ham_pdf, dpi=200)
    except Exception as e:
        raise RuntimeError("Poppler (pdf2image) yüklenmemiş veya PDF işlenemedi.") from e

    parcalar = []
    for sayfa in sayfalar:
        try:
            metin = pytesseract.image_to_string(sayfa, lang="tur+eng")
            if metin:
                parcalar.append(metin.strip())
        except Exception as e:
            raise RuntimeError("Tesseract-OCR yüklenmemiş veya metin çıkarılamadı.") from e
            
    return "\n\n".join(parcalar).encode("utf-8")


def docx_bytes_to_txt(ham_docx: bytes) -> bytes:
    """DOCX paragraflarını birleştirir."""
    import docx

    belge = docx.Document(io.BytesIO(ham_docx))
    satirlar = [p.text for p in belge.paragraphs if p.text.strip()]
    return "\n".join(satirlar).encode("utf-8")


def pptx_bytes_to_txt(ham_pptx: bytes) -> bytes:
    """PPTX slayt metinlerini çıkarır."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(ham_pptx))
    satirlar: list[str] = []
    for slayt in prs.slides:
        for sekil in slayt.shapes:
            if hasattr(sekil, "text") and sekil.text and sekil.text.strip():
                satirlar.append(sekil.text.strip())
    return "\n".join(satirlar).encode("utf-8")


def metin_bytes_normalize(ham: bytes) -> bytes:
    """UTF-8 BOM kaldırır ve satır sonlarını LF yapar."""
    t = ham.decode("utf-8-sig")
    t = t.replace("\r\n", "\n").replace("\r", "\n")
    return t.encode("utf-8")


def txt_cikis_dosya_adi(orijinal_ad: str) -> str:
    kok = PurePath(orijinal_ad or "cikti").stem
    return f"{kok}.txt"


def json_cikis_dosya_adi(orijinal_ad: str) -> str:
    """İndirilecek JSON dosya adı (uzantı .json)."""
    kok = PurePath(orijinal_ad or "cikti").stem
    return f"{kok}.json"


def tsv_cikis_dosya_adi(orijinal_ad: str) -> str:
    """CSV kaynağından TSV çıktı adı."""
    kok = PurePath(orijinal_ad or "cikti").stem
    return f"{kok}.tsv"


def _rgb_beyaz_arkaplan(im: Image.Image) -> Image.Image:
    """Şeffaf görüntüyü beyaz fona yapıştırır (BMP/JPEG uyumu için)."""
    rgba = im.convert("RGBA")
    beyaz = Image.new("RGB", rgba.size, (255, 255, 255))
    beyaz.paste(rgba, mask=rgba.split()[3])
    return beyaz


def resim_bytes_donustur(ham_resim: bytes, hedef_format: str) -> tuple[bytes, str]:
    """
    Pillow ile görüntüyü açıp istenen formatta bayta yazar.

    hedef_format: PNG, JPEG, WEBP, BMP, TIFF, GIF (büyük harf).
    Dönüş: (çıkış baytları, uzantı — küçük harf, noktasız).
    """
    fmt = hedef_format.strip().upper()
    if fmt not in {"PNG", "JPEG", "WEBP", "BMP", "TIFF", "GIF"}:
        raise ValueError("desteklenmeyen_resim_formati")

    im = Image.open(io.BytesIO(ham_resim))
    cikis = io.BytesIO()
    kayit_arg: dict = {"format": fmt}

    if fmt == "JPEG":
        if im.mode in ("RGBA", "LA", "P"):
            im = _rgb_beyaz_arkaplan(im if im.mode != "P" else im.convert("RGBA"))
        elif im.mode != "RGB":
            im = im.convert("RGB")
        kayit_arg["quality"] = 90
    elif fmt == "BMP":
        if im.mode in ("RGBA", "LA", "P"):
            im = _rgb_beyaz_arkaplan(im.convert("RGBA") if im.mode == "P" else im)
        else:
            im = im.convert("RGB")
    elif fmt == "GIF":
        # Animasyonlu GIF'lerde yalnızca ilk kare işlenir
        if im.mode in ("RGBA", "LA"):
            im = im.convert("RGBA")
            im = im.convert("P", palette=Image.ADAPTIVE, colors=256)
        elif im.mode not in ("P", "RGB", "L"):
            im = im.convert("RGB").convert("P", palette=Image.ADAPTIVE, colors=256)
    elif fmt == "TIFF":
        if im.mode not in ("RGB", "RGBA", "L", "P"):
            im = im.convert("RGB")
        kayit_arg["compression"] = "tiff_lzw"

    im.save(cikis, **kayit_arg)
    uzanti = {"PNG": "png", "JPEG": "jpg", "WEBP": "webp", "BMP": "bmp", "TIFF": "tiff", "GIF": "gif"}[fmt]
    return cikis.getvalue(), uzanti


def resim_cikis_dosya_adi(orijinal_ad: str, uzanti: str) -> str:
    """İndirilecek resim adı."""
    kok = PurePath(orijinal_ad or "cikti").stem
    return f"{kok}.{uzanti}"
